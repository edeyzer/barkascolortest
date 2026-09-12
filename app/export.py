# app/export.py - Katalogni Excel/PDF formatida eksport qilish
import os
import tempfile
from io import BytesIO

from telegram import Update, InlineKeyboardButton, InlineKeyboardMarkup
from telegram.ext import ContextTypes
from PIL import Image as PILImage

from app.admin import is_admin
from app.database import fetchall
from app.utils import logger

THUMB_SIZE = 160  # eksportdagi rasm o'lchami (piksel)
PROGRESS_STEP = 20  # necha kodda bir marta progress xabari yuborilsin


async def export_cmd(update: Update, context: ContextTypes.DEFAULT_TYPE):
    """/export - Adminlar uchun: katalogni Excel yoki PDF formatida eksport qilish."""
    user_id = update.effective_user.id
    if not await is_admin(user_id):
        await update.message.reply_text("❌ Faqat adminlar buyruqdan foydalanishi mumkin.")
        return

    buttons = [[
        InlineKeyboardButton("📊 Excel", callback_data="export_excel"),
        InlineKeyboardButton("📄 PDF", callback_data="export_pdf"),
    ]]
    await update.message.reply_text(
        "📥 Katalogni qaysi formatda eksport qilaylik?",
        reply_markup=InlineKeyboardMarkup(buttons)
    )


async def _download_thumb(bot, file_id: str):
    """Telegram file_id orqali rasmni yuklab, kichraytirib, vaqtinchalik faylga saqlaydi.
    Fayl yo'lini qaytaradi, muvaffaqiyatsiz bo'lsa None."""
    if not file_id:
        return None
    try:
        file = await bot.get_file(file_id)
        bio = BytesIO()
        await file.download_to_memory(bio)
        bio.seek(0)
        img = PILImage.open(bio).convert("RGB")
        img.thumbnail((THUMB_SIZE, THUMB_SIZE))
        tmp = tempfile.NamedTemporaryFile(suffix=".jpg", delete=False)
        img.save(tmp.name, format="JPEG", quality=85)
        tmp.close()
        return tmp.name
    except Exception as e:
        logger.error(f"Eksport uchun rasm yuklashda xato: {e}")
        return None


async def handle_export(update: Update, context: ContextTypes.DEFAULT_TYPE):
    """export_excel / export_pdf callback tugmalarini boshqaradi."""
    query = update.callback_query
    user_id = query.from_user.id

    if not await is_admin(user_id):
        await query.answer("❌ Ruxsat yo'q", show_alert=True)
        return

    fmt = "excel" if query.data == "export_excel" else "pdf"
    await query.answer()
    await query.message.edit_text("🔄 Eksport tayyorlanmoqda... Bu biroz vaqt olishi mumkin.")

    rows = await fetchall(
        "SELECT code, file_id, dominant_hex, color_name, search_count, album FROM images ORDER BY code"
    )

    if not rows:
        await query.message.edit_text("📭 Katalogda hozircha hech qanday kod yo'q.")
        return

    items = []
    for row in rows:
        if isinstance(row, tuple):
            items.append({
                "code": row[0], "file_id": row[1], "hex": row[2],
                "name": row[3], "count": row[4] or 0, "album": row[5] or "",
            })
        else:
            items.append({
                "code": row.get("code"), "file_id": row.get("file_id"),
                "hex": row.get("dominant_hex"), "name": row.get("color_name"),
                "count": row.get("search_count") or 0, "album": row.get("album") or "",
            })

    temp_paths = []
    filepath = None
    try:
        total = len(items)
        for i, item in enumerate(items, 1):
            item["thumb_path"] = await _download_thumb(context.bot, item["file_id"])
            if item["thumb_path"]:
                temp_paths.append(item["thumb_path"])

            if total > PROGRESS_STEP and i % PROGRESS_STEP == 0:
                try:
                    await query.message.edit_text(f"🔄 Rasmlar yuklanmoqda: {i}/{total}")
                except Exception:
                    pass

        filepath = _build_excel(items) if fmt == "excel" else _build_pdf(items)

        with open(filepath, "rb") as f:
            await query.message.reply_document(
                document=f,
                filename=os.path.basename(filepath),
                caption=f"✅ Katalog eksport qilindi!\n📦 Jami: {total} ta kod"
            )
        await query.message.delete()

    except Exception as e:
        logger.error(f"Eksport xatosi: {e}")
        try:
            await query.message.edit_text(f"❌ Eksport qilishda xatolik: {e}")
        except Exception:
            pass
    finally:
        for p in temp_paths:
            try:
                os.remove(p)
            except Exception:
                pass
        if filepath:
            try:
                os.remove(filepath)
            except Exception:
                pass


def _build_excel(items):
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.styles import Font, Alignment

    wb = Workbook()
    ws = wb.active
    ws.title = "Katalog"

    headers = ["#", "Rasm", "Kod", "HEX", "Rang nomi", "Qidirilgan", "Albom"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")

    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 20
    ws.column_dimensions["C"].width = 14
    ws.column_dimensions["D"].width = 12
    ws.column_dimensions["E"].width = 22
    ws.column_dimensions["F"].width = 12
    ws.column_dimensions["G"].width = 16

    for idx, item in enumerate(items, start=2):
        ws.row_dimensions[idx].height = 100
        ws.cell(row=idx, column=1, value=idx - 1)
        ws.cell(row=idx, column=3, value=item["code"])
        ws.cell(row=idx, column=4, value=item["hex"] or "")
        ws.cell(row=idx, column=5, value=item["name"] or "")
        ws.cell(row=idx, column=6, value=item["count"])
        ws.cell(row=idx, column=7, value=item["album"])

        if item.get("thumb_path"):
            try:
                img = XLImage(item["thumb_path"])
                img.width = 120
                img.height = 120
                ws.add_image(img, f"B{idx}")
            except Exception as e:
                logger.error(f"Excel rasm qo'shishda xato ({item['code']}): {e}")

    tmp = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
    tmp.close()
    wb.save(tmp.name)
    return tmp.name


def _build_pdf(items):
    from reportlab.lib.pagesizes import A4
    from reportlab.lib import colors
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Image as RLImage, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet

    tmp = tempfile.NamedTemporaryFile(suffix=".pdf", delete=False)
    tmp.close()

    doc = SimpleDocTemplate(
        tmp.name, pagesize=A4,
        topMargin=15 * mm, bottomMargin=15 * mm,
        leftMargin=12 * mm, rightMargin=12 * mm
    )
    styles = getSampleStyleSheet()

    data = [["Rasm", "Kod", "HEX", "Rang nomi", "Qidirilgan"]]
    for item in items:
        img_cell = ""
        if item.get("thumb_path"):
            try:
                img_cell = RLImage(item["thumb_path"], width=22 * mm, height=22 * mm)
            except Exception:
                img_cell = ""
        data.append([
            img_cell,
            Paragraph(item["code"] or "", styles["Normal"]),
            Paragraph(item["hex"] or "", styles["Normal"]),
            Paragraph(item["name"] or "", styles["Normal"]),
            str(item["count"]),
        ])

    table = Table(
        data,
        colWidths=[26 * mm, 26 * mm, 22 * mm, 55 * mm, 25 * mm],
        repeatRows=1
    )
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#333333")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f5f5f5")]),
    ]))

    story = [
        Paragraph("BarkasColor - Ranglar katalogi", styles["Title"]),
        Spacer(1, 6 * mm),
        table,
    ]
    doc.build(story)
    return tmp.name